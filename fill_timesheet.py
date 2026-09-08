"""
fill_timesheet.py
==================

Remplit le fichier de suivi de formation (xlsx, feuille "DASSOULI
Zéphyr") pour une semaine donnée, à partir de l'EDT IMT Atlantique.
Réutilise ics_calendar.py (doit être dans le même dossier) pour la
récupération/le parsing de l'ICS et la classification en code UE.

Zones remplies dans le classeur (conventions du fichier fourni) :
    C10 : date de début de semaine ("Semaine du")
    E10 : date de fin de semaine ("au")
    G10 : numéro de semaine ISO ("N° de semaine")
    lignes 22 à 49, une par séance :
        B : date (dd/mm)
        C : code UE
        D : horaire début
        E : horaire fin
        F : durée de pause (min)
        G : durée de formation (h, affichée "HH"mm)
    J10 : total des heures de formation de la semaine (formule Excel
          =SUM(...), recalculée automatiquement -> jamais une valeur
          codée en dur)

Fusion des créneaux
--------------------
Le flux PASS peut découper une séance continue en plusieurs VEVENT
(ex: cours + pause + TD). Deux créneaux consécutifs (dans l'ordre
chronologique) qui partagent le même code UE ET tombent le même jour
sont donc fusionnés en une seule ligne :
    - horaire début = début du premier créneau du groupe
    - horaire fin   = fin du dernier créneau du groupe
    - durée de pause = temps "creux" entre les créneaux du groupe
                        (somme des trous), en minutes
    - durée de formation = somme des durées réelles de chaque créneau
                        (= temps total du groupe moins la pause)
Deux créneaux de même code mais de jours différents ne sont PAS
fusionnés (ça n'aurait aucun sens de "fusionner" à travers plusieurs
jours) : c'est une hypothèse de bon sens, pas une consigne explicite,
donc à signaler si un cas réel s'avère différent.
Une pause de 60 minutes ou plus sépare également les créneaux en deux
séances distinctes.

Usage
-----
    export ICS_CALENDAR_URL="https://inpass.imt-atlantique.fr/passcal/getics?login=...&check=..."
    python3 fill_timesheet.py --week 37 --input DASSOULI_Zephyr.xlsx

    # semaine à cheval sur deux années scolaires -> lever l'ambiguïté
    python3 fill_timesheet.py --week 36 --year 2026 --input DASSOULI_Zephyr.xlsx
"""

from __future__ import annotations

import argparse
import logging
from dataclasses import dataclass
from datetime import date, time, timedelta
from pathlib import Path
from typing import Optional

import openpyxl

from ics_calendar import Event, classify_ue, get_schedule, DEFAULT_EXTRACTION_START

logger = logging.getLogger("fill_timesheet")
if not logger.handlers:
    _h = logging.StreamHandler()
    _h.setFormatter(logging.Formatter("[%(levelname)s] %(message)s"))
    logger.addHandler(_h)
logger.setLevel(logging.INFO)

SHEET_NAME = "DASSOULI Zéphyr"
FIRST_DATA_ROW = 22
LAST_DATA_ROW = 49  # étendue du tableau bordé dans le modèle fourni
DATE_NUMBER_FORMAT = "dd/mm"
TIME_NUMBER_FORMAT = "h:mm"
# Format "élapsé" (les crochets empêchent Excel de faire un modulo 24h,
# indispensable pour le total J10 qui peut dépasser 24h) avec un "H"
# littéral comme séparateur, ex: 2H30.
DURATION_NUMBER_FORMAT = '[h]"H"mm'
MAX_MERGED_PAUSE = timedelta(minutes=59)


@dataclass
class Session:
    """Une ligne du tableau, après fusion des créneaux consécutifs de même code UE."""

    day: date
    code_ue: str
    start: time
    end: time
    pause_minutes: int
    formation: timedelta  # durée nette de formation (pause déjà déduite)


# --------------------------------------------------------------------------
# Fusion des créneaux consécutifs
# --------------------------------------------------------------------------

def merge_sessions(events: list[Event]) -> list[Session]:
    events = sorted(events, key=lambda e: e.start)
    sessions: list[Session] = []
    group: list[Event] = []

    def flush(grp: list[Event]) -> None:
        if not grp:
            return
        first, last = grp[0], grp[-1]
        span = last.end - first.start
        total_active = sum((e.end - e.start for e in grp), timedelta())
        pause = span - total_active
        sessions.append(
            Session(
                day=first.start.date(),
                code_ue=classify_ue(first.summary),
                start=first.start.time(),
                end=last.end.time(),
                pause_minutes=round(pause.total_seconds() / 60),
                formation=total_active,
            )
        )

    for e in events:
        if e.end is None:
            logger.warning("Évènement sans heure de fin ignoré (uid=%s): %s", e.uid, e.summary)
            continue
        if not group:
            group = [e]
            continue
        same_day = e.start.date() == group[-1].start.date()
        same_code = classify_ue(e.summary) == classify_ue(group[-1].summary)
        pause = e.start - group[-1].end
        if same_day and same_code and pause <= MAX_MERGED_PAUSE:
            group.append(e)
        else:
            flush(group)
            group = [e]
    flush(group)

    return sessions


# --------------------------------------------------------------------------
# Sélection de la semaine
# --------------------------------------------------------------------------

def select_week_events(
    all_events: list[Event], week: int, year: Optional[int] = None
) -> tuple[list[Event], int]:
    """
    Filtre les évènements dont la semaine ISO == `week`.

    Si `year` n'est pas précisé et que plusieurs années correspondent
    (cas limite en tout début/fin d'année scolaire, ex: semaine 36
    présente à la fois en septembre 2026 et septembre 2027), la plus
    ancienne est retenue et un avertissement est émis -- repasser
    `--year` pour lever l'ambiguïté si ce n'est pas la bonne.
    """
    candidates = [e for e in all_events if e.start.isocalendar()[1] == week]
    years_found = sorted({e.start.isocalendar()[0] for e in candidates})

    if not years_found:
        raise ValueError(f"Aucun évènement trouvé pour la semaine ISO {week}.")

    if year is not None:
        chosen_year = year
    elif len(years_found) > 1:
        chosen_year = years_found[0]
        logger.warning(
            "La semaine %d existe sur plusieurs années %s : utilisation de %d "
            "(passe --year pour choisir explicitement).",
            week, years_found, chosen_year,
        )
    else:
        chosen_year = years_found[0]

    filtered = [e for e in candidates if e.start.isocalendar()[0] == chosen_year]
    if not filtered:
        raise ValueError(f"Aucun évènement pour la semaine ISO {week} de l'année {chosen_year}.")
    return filtered, chosen_year


# --------------------------------------------------------------------------
# Écriture dans le classeur
# --------------------------------------------------------------------------

def fill_workbook(
    input_path: Path,
    output_path: Path,
    sessions: list[Session],
    week: int,
    year: int,
) -> None:
    wb = openpyxl.load_workbook(input_path)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(
            f"Feuille '{SHEET_NAME}' introuvable dans {input_path} (feuilles: {wb.sheetnames})"
        )
    ws = wb[SHEET_NAME]

    monday = date.fromisocalendar(year, week, 1)
    sunday = date.fromisocalendar(year, week, 7)

    # --- en-tête de semaine ---
    ws["C10"] = monday
    ws["C10"].number_format = DATE_NUMBER_FORMAT
    ws["E10"] = sunday
    ws["E10"].number_format = DATE_NUMBER_FORMAT
    ws["G10"] = week

    # --- purge du tableau avant d'écrire la nouvelle semaine (évite de
    #     laisser trainer des lignes d'un run précédent sur une autre
    #     semaine) ---
    for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        for col in "BCDEFG":
            ws[f"{col}{row}"] = None

    n_available = LAST_DATA_ROW - FIRST_DATA_ROW + 1
    if len(sessions) > n_available:
        logger.warning(
            "%d séance(s) à écrire mais seulement %d ligne(s) disponibles (%d-%d) : "
            "les %d dernière(s) séance(s) ne seront pas écrites.",
            len(sessions), n_available, FIRST_DATA_ROW, LAST_DATA_ROW,
            len(sessions) - n_available,
        )

    for i, s in enumerate(sessions[:n_available]):
        r = FIRST_DATA_ROW + i
        b = ws[f"B{r}"]
        b.value = s.day
        b.number_format = DATE_NUMBER_FORMAT
        ws[f"C{r}"] = s.code_ue
        d = ws[f"D{r}"]
        d.value = s.start
        d.number_format = TIME_NUMBER_FORMAT
        e = ws[f"E{r}"]
        e.value = s.end
        e.number_format = TIME_NUMBER_FORMAT
        ws[f"F{r}"] = s.pause_minutes
        g = ws[f"G{r}"]
        g.value = s.formation
        g.number_format = DURATION_NUMBER_FORMAT

    # --- total : formule Excel, jamais une valeur codée en dur ---
    j10 = ws["J10"]
    j10.value = f"=SUM(G{FIRST_DATA_ROW}:G{LAST_DATA_ROW})"
    j10.number_format = DURATION_NUMBER_FORMAT

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    logger.info(
        "Classeur écrit: %s (%d séance(s))", output_path, min(len(sessions), n_available)
    )


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

def _build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Remplit le suivi de formation xlsx pour une semaine donnée."
    )
    p.add_argument("--week", type=int, required=True, help="Numéro de semaine ISO (1-53)")
    p.add_argument("--year", type=int, default=None, help="Année ISO (désambiguïsation si besoin)")
    p.add_argument("--input", type=Path, required=True, help="Fichier xlsx modèle à remplir")
    p.add_argument(
        "--output", type=Path, default=None,
        help="Fichier de sortie (défaut: <input>_S<semaine>.xlsx, à côté de l'input)",
    )
    p.add_argument("--url", help="URL du flux ICS (sinon: variable d'env ICS_CALENDAR_URL)")
    p.add_argument("--cache", default="schedule_cache.ics", help="Chemin du cache local ICS")
    p.add_argument("--cache-max-age", type=int, default=900, help="Âge max du cache en secondes")
    p.add_argument("--force-refresh", action="store_true", help="Ignore le cache et retélécharge")
    p.add_argument("-v", "--verbose", action="store_true")
    return p


def main(argv: Optional[list[str]] = None) -> int:
    args = _build_arg_parser().parse_args(argv)
    if args.verbose:
        logger.setLevel(logging.DEBUG)

    try:
        all_events = get_schedule(
            url=args.url,
            cache_path=args.cache,
            cache_max_age=args.cache_max_age,
            force_refresh=args.force_refresh,
            since=DEFAULT_EXTRACTION_START,
        )
        week_events, year = select_week_events(all_events, args.week, args.year)
        sessions = merge_sessions(week_events)
    except Exception as exc:
        logger.error("Échec: %s", exc)
        return 1

    output_path = args.output or args.input.with_name(f"{args.input.stem}_S{args.week:02d}.xlsx")

    try:
        fill_workbook(args.input, output_path, sessions, args.week, year)
    except Exception as exc:
        logger.error("Échec lors de l'écriture du classeur: %s", exc)
        return 1

    total = sum((s.formation for s in sessions), timedelta())
    print(f"\nSemaine ISO {args.week} ({year}) : {len(sessions)} séance(s), total formation = {total}")
    for s in sessions:
        print(
            f"  {s.day:%d/%m}  {s.code_ue:<10} {s.start:%H:%M}-{s.end:%H:%M}"
            f"  pause={s.pause_minutes:>3}min  formation={s.formation}"
        )

    return 0


if __name__ == "__main__":
    raise SystemExit(main())